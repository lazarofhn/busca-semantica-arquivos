"""
extrair.py — varre a pasta de trabalho, extrai texto e gera trechos (chunks) p/ embedding.

Saídas (em dados/):
  chunks.jsonl.gz   1 linha por trecho: {id, arq, pi, pf, text}
  arquivos.jsonl    1 linha por arquivo processado: {arq, mtime, size, ext, status, paginas, pag_sem_texto, n_chunks, dup_de}
Retomável: arquivos já em arquivos.jsonl (mesmo mtime/size) são pulados.
Uso: python extrair.py [--limite N]
"""
import gzip
import hashlib
import json
import os
import sys
import time
from multiprocessing import Pool

import config

RAIZ = config.RAIZ                 # pasta indexada (config.json → "raiz")
DADOS = config.DADOS
EXTS = {".pdf", ".docx", ".xlsx", ".md", ".txt"}
PULAR_DIRS = config.PASTAS_IGNORADAS
MAX_TXT = 5 * 1024 * 1024          # .txt maior que isso costuma ser log/dado bruto
MAX_CHARS_ARQ = 3_000_000          # teto de texto por arquivo (planilha/log gigante)
CHUNK = 1500                       # ~400 tokens
SOBRA = 200                        # sobreposição entre trechos
MIN_PAG = 40                       # página com menos que isso = provável digitalização


def listar():
    for d, subdirs, files in os.walk(RAIZ):
        subdirs[:] = [s for s in subdirs if s not in PULAR_DIRS and not s.startswith(".")]
        for f in files:
            if f.startswith("~$"):
                continue
            ext = os.path.splitext(f)[1].lower()
            if ext in EXTS:
                yield os.path.join(d, f)


def paginas_pdf(p):
    import fitz
    with fitz.open(p) as doc:
        return [pg.get_text("text") for pg in doc]


def paginas_docx(p):
    import docx
    d = docx.Document(p)
    partes = [par.text for par in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            partes.append(" | ".join(c.text.strip() for c in r.cells))
    return ["\n".join(partes)]


MAX_CHARS_XLSX = 200_000           # planilha de notas/lançamentos: o começo basta p/ achar o arquivo


class DadosBrutos(Exception):
    """Arquivo de dados (SPED/ECF, exportação) — não é texto para busca."""


def paginas_xlsx(p):
    from python_calamine import CalamineWorkbook
    wb = CalamineWorkbook.from_path(p)
    out, tot = [], 0
    for nome in wb.sheet_names:
        linhas = [f"[Planilha: {nome}]"]
        try:
            rows = wb.get_sheet_by_name(nome).to_python(skip_empty_area=True)
        except Exception:
            continue
        for row in rows:
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals:
                linhas.append(" | ".join(vals))
                tot += len(linhas[-1])
            if tot > MAX_CHARS_XLSX:
                break
        out.append("\n".join(linhas))
        if tot > MAX_CHARS_XLSX:
            break
    return out


def parece_dados(texto):
    linhas = [l for l in texto[:200_000].split("\n") if l.strip()]
    if len(linhas) < 50:
        return False
    pipes = sum(1 for l in linhas if l.lstrip().startswith("|")) / len(linhas)      # SPED/ECF/EFD
    amostra = texto[:200_000]
    letras = sum(c.isalpha() for c in amostra) / max(1, len(amostra))
    return pipes > 0.3 or letras < 0.35


def paginas_txt(p):
    if os.path.getsize(p) > MAX_TXT:
        raise ValueError("txt grande demais")
    raw = open(p, "rb").read()
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            t = raw.decode(enc)
            break
        except UnicodeDecodeError:
            pass
    else:
        t = raw.decode("utf-8", "ignore")
    if parece_dados(t):
        raise DadosBrutos()
    return [t]


LEITORES = {".pdf": paginas_pdf, ".docx": paginas_docx, ".xlsx": paginas_xlsx,
            ".md": paginas_txt, ".txt": paginas_txt}


def limpar(s):
    s = s.replace("\x00", " ").replace("\r", "")
    linhas = [" ".join(l.split()) for l in s.split("\n")]
    out, vazio = [], False
    for l in linhas:
        if l:
            out.append(l); vazio = False
        elif not vazio:
            out.append(""); vazio = True
    return "\n".join(out).strip()


def fatiar(paginas):
    """Concatena páginas e corta em trechos ~CHUNK chars, cortando em espaço; devolve (texto, pag_ini, pag_fim)."""
    buf, mapa = [], []          # mapa: (offset_inicio, n_pagina)
    pos = 0
    for i, t in enumerate(paginas, 1):
        if not t:
            continue
        mapa.append((pos, i))
        buf.append(t)
        pos += len(t) + 2
    texto = "\n\n".join(buf)
    if not texto:
        return []

    def pag(off):
        n = mapa[0][1]
        for o, k in mapa:
            if o <= off:
                n = k
            else:
                break
        return n

    out, ini, N = [], 0, len(texto)
    while ini < N:
        fim = min(ini + CHUNK, N)
        if fim < N:
            corte = max(texto.rfind("\n", ini + CHUNK // 2, fim), texto.rfind(" ", ini + CHUNK // 2, fim))
            if corte > ini:
                fim = corte
        trecho = texto[ini:fim].strip()
        if len(trecho) >= 30:
            out.append((trecho, pag(ini), pag(max(ini, fim - 1))))
        if fim >= N:
            break
        ini = max(fim - SOBRA, ini + 1)
    return out


def processar(p):
    rel = os.path.relpath(p, RAIZ)
    ext = os.path.splitext(p)[1].lower()
    st = os.stat(p)
    reg = {"arq": rel, "mtime": int(st.st_mtime), "size": st.st_size, "ext": ext}
    try:
        pags = [limpar(t) for t in LEITORES[ext](p)]
    except DadosBrutos:
        reg.update(status="dados_brutos")
        return reg, []
    except Exception as e:
        reg.update(status="erro", erro=repr(e)[:200])
        return reg, []
    if sum(len(t) for t in pags) > MAX_CHARS_ARQ:
        acc, cort = 0, []
        for t in pags:
            if acc > MAX_CHARS_ARQ:
                break
            cort.append(t); acc += len(t)
        pags = cort
        reg["truncado"] = True
    reg["paginas"] = len(pags)
    if ext == ".pdf":
        reg["pag_sem_texto"] = sum(1 for t in pags if len(t) < MIN_PAG)
    trechos = fatiar(pags)
    reg["hash"] = hashlib.sha1("\n".join(pags).encode("utf-8", "ignore")).hexdigest()
    reg["status"] = "ok" if trechos else "sem_texto"
    reg["n_chunks"] = len(trechos)
    return reg, trechos


def main():
    config.exigir_raiz()
    limite = int(sys.argv[sys.argv.index("--limite") + 1]) if "--limite" in sys.argv else None
    os.makedirs(DADOS, exist_ok=True)
    f_arq = os.path.join(DADOS, "arquivos.jsonl")
    f_chk = os.path.join(DADOS, "chunks.jsonl.gz")

    feitos, hashes, prox_id = {}, {}, 0
    if os.path.exists(f_arq):
        for l in open(f_arq, encoding="utf-8"):
            r = json.loads(l)
            feitos[r["arq"]] = (r["mtime"], r["size"])
            if r.get("hash") and not r.get("dup_de"):
                hashes.setdefault(r["hash"], r["arq"])
            prox_id = max(prox_id, r.get("id_fim", -1) + 1)

    todos = list(listar())
    pend = [p for p in todos
            if feitos.get(os.path.relpath(p, RAIZ)) != (int(os.stat(p).st_mtime), os.stat(p).st_size)]
    print(f"{len(todos)} arquivos elegíveis, {len(todos) - len(pend)} já feitos, {len(pend)} pendentes", flush=True)
    if limite:
        pend = pend[:limite]

    t0, n, nchunks = time.time(), 0, 0
    with Pool(4, maxtasksperchild=200) as pool, \
            open(f_arq, "a", encoding="utf-8") as fa, \
            gzip.open(f_chk, "at", encoding="utf-8", compresslevel=6) as fc:
        for reg, trechos in pool.imap_unordered(processar, pend, chunksize=4):
            n += 1
            h = reg.get("hash")
            if h and trechos and h in hashes and hashes[h] != reg["arq"]:
                reg["dup_de"] = hashes[h]          # mesmo texto já indexado em outro arquivo
                trechos = []
            elif h and trechos:
                hashes[h] = reg["arq"]
            if trechos:
                reg["id_ini"] = prox_id
                for t, pi, pf in trechos:
                    fc.write(json.dumps({"id": prox_id, "arq": reg["arq"], "pi": pi, "pf": pf, "text": t},
                                        ensure_ascii=False) + "\n")
                    prox_id += 1
                reg["id_fim"] = prox_id - 1
                nchunks += len(trechos)
            fa.write(json.dumps(reg, ensure_ascii=False) + "\n")
            if n % 200 == 0:
                fa.flush(); fc.flush()
                el = time.time() - t0
                print(f"[{n}/{len(pend)}] {nchunks} trechos | {n/el:.1f} arq/s | "
                      f"faltam ~{(len(pend)-n)/(n/el)/60:.0f} min", flush=True)
    print(f"FIM: {n} arquivos, {nchunks} trechos novos, {time.time()-t0:.0f}s", flush=True)


if __name__ == "__main__":
    main()
